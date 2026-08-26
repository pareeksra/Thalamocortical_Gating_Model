import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
import os

def create_presentation():
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    prs = Presentation()
    
    # Define slide layouts
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]
    title_only_layout = prs.slide_layouts[5]

    # --- 1. Title Slide ---
    slide = prs.slides.add_slide(title_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Thalamocortical Gating Project"
    subtitle.text = "Model Progress, Phenomenological vs Mechanistic Approaches, and New Protocols"

    # --- 2. The Emergent Mechanistic Model ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "1. The Emergent Mechanistic Model"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(1.5))
    tf = txBox.text_frame
    tf.text = "Our New Mathematical Architecture:"
    p = tf.add_paragraph()
    p.text = "• Inhibition is dynamically calculated using firing rates: MD Thalamus -> VIP -> SOM."
    p = tf.add_paragraph()
    p.text = "• VIP successfully suppresses SOM, dropping dendritic inhibition organically to ~2 nS."
    
    img_path = os.path.join("figures", "mechanistic", "fig_emergent_inhibition.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.5), Inches(3.0), height=Inches(4.0))

    # --- 3. The 3-Epoch Protocol ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "2. The New 3-Epoch Protocol"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9.0), Inches(1.5))
    tf = txBox.text_frame
    tf.text = "Continuous 300-Trial Simulation Paradigm:"
    p = tf.add_paragraph()
    p.text = "• Baseline (0-99): No MD drive, SOM inhibition at maximum."
    p = tf.add_paragraph()
    p.text = "• Context A (100-199): MD activates VIP 1. Branch 1 learns."
    p = tf.add_paragraph()
    p.text = "• Context B (200-299): MD activates VIP 2. Branch 2 learns."
    
    img_path = os.path.join("figures", "mechanistic", "fig_self_organized_branching.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.5), Inches(3.2), height=Inches(4.0))

    # --- 4. Gating Selectivity Results ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "3. Gating Selectivity Index (GS)"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9.0), Inches(1.5))
    tf = txBox.text_frame
    tf.text = "Quantifying interneuron behavior:"
    p = tf.add_paragraph()
    p.text = "• VIP GS = 1.00 (Highly selective for preferred context)"
    p = tf.add_paragraph()
    p.text = "• SOM GS = -0.88 (Anti-selective, strongly suppressed by VIP)"
    p = tf.add_paragraph()
    p.text = "• PV GS = 0.00 (Global, non-selective drive)"
    
    img_path = os.path.join("figures", "mechanistic", "fig_gating_selectivity.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.5), Inches(3.2), width=Inches(7.0))

    # --- 5. The Phenomenological Model: Calcium Plasticity ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "4. Phenomenological Model: Calcium Plasticity"
    
    img_path = os.path.join("figures", "phenomenological", "fig1_calcium_plasticity.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.0), Inches(1.5), width=Inches(8.0))

    # --- 6. The Phenomenological Model: Dendritic Gating ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "5. Phenomenological Model: Dendritic Gating"
    
    img_path = os.path.join("figures", "phenomenological", "fig2_dendritic_gating.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.0), Inches(1.5), width=Inches(8.0))

    # --- 7. The Phenomenological Model: Inhibition vs Plasticity ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "6. Phenomenological Model: Inhibition vs Plasticity"
    
    img_path = os.path.join("figures", "phenomenological", "fig3_weight_change_vs_rate.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(2.0), Inches(1.5), height=Inches(5.0))

    # --- 8. The Phenomenological Model: Branch-Specific Learning ---
    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "7. Phenomenological Model: Branch-Specific Learning"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9.0), Inches(1.5))
    tf = txBox.text_frame
    tf.text = "Original Approach (Yang et al. 2016):"
    p = tf.add_paragraph()
    p.text = "• Uses hardcoded inhibition values (e.g., 2 nS vs 35 nS) to force gating."
    p = tf.add_paragraph()
    p.text = "• Proves the mathematical concept of Graupner-Brunel plasticity."
    
    img_path = os.path.join("figures", "phenomenological", "fig4_branch_specific_learning.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1.0), Inches(3.0), width=Inches(8.0))

    # --- 9. Future Directions ---
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "8. Future Directions"
    tf = slide.placeholders[1].text_frame
    
    p = tf.paragraphs[0]
    p.text = "Inject Biological Realism:"
    
    p = tf.add_paragraph()
    p.text = "Introduce stochastic variance (Gaussian noise) to firing rates and synaptic weights."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Signal-to-Noise Ratio (SNR) Paradigm:"
    
    p = tf.add_paragraph()
    p.text = "Test the circuit's ability to fork pathways in high-noise vs low-noise surroundings."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Simulate Ketamine/NMDA Hypofunction:"
    
    p = tf.add_paragraph()
    p.text = "Degrade MD->VIP synapses and observe emergent gating failure (Schizophrenia model)."
    p.level = 1

    prs.save(os.path.join("docs", "presentations", "Thalamocortical_Gating_Project_Summary_v3.pptx"))
    print("Presentation created successfully as 'Thalamocortical_Gating_Project_Summary_v3.pptx'")

if __name__ == "__main__":
    create_presentation()
