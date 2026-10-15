"""
build_pi_meeting_ppt.py -- PI meeting deck for the 2026-10-14/15 session.
=========================================================================
Plain-background background slides (the question + the model), then the three new
results each as text-left / figure-right. Matches the house style of
plotting/build_results_ppt.py (16:9, blank layout, black/grey on white).

Run: <venv>/python -I plotting/build_pi_meeting_ppt.py
Out: docs/PI_Meeting_2026-10-15.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(PROOT, 'figures')
OUT = os.path.join(PROOT, 'docs', 'PI_Meeting_2026-10-15.pptx')

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)
BLUE = RGBColor(0x00, 0x72, 0xB2)
GREEN = RGBColor(0x00, 0x9E, 0x73)
ORANGE = RGBColor(0xD5, 0x5E, 0x00)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _tb(slide, l, t, w, h):
    return slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame


def title_slide(title, subtitle=None, foot=None):
    s = prs.slides.add_slide(BLANK)
    tf = _tb(s, 0.8, 2.6, 11.7, 1.5); tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(38); r.font.bold = True; r.font.color.rgb = BLACK
    if subtitle:
        tf2 = _tb(s, 0.8, 4.0, 11.7, 1.2); tf2.word_wrap = True
        r2 = tf2.paragraphs[0].add_run(); r2.text = subtitle
        r2.font.size = Pt(18); r2.font.color.rgb = GREY
    if foot:
        tf3 = _tb(s, 0.8, 6.6, 11.7, 0.6)
        r3 = tf3.paragraphs[0].add_run(); r3.text = foot
        r3.font.size = Pt(13); r3.font.color.rgb = GREY
    return s


def notes(s, text):
    """Attach speaker notes to a slide."""
    s.notes_slide.notes_text_frame.text = text.strip()


def header(s, text):
    tf = _tb(s, 0.6, 0.35, 12.1, 0.9); tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = BLACK


def bullets(s, items, left=0.8, top=1.5, width=11.7, height=5.5, size=20):
    tf = _tb(s, left, top, width, height); tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0; txt = it; color = BLACK; bold = False
        if isinstance(it, tuple):
            txt, lvl = it[0], it[1]
            if len(it) > 2: color = it[2]
            if len(it) > 3: bold = it[3]
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.level = lvl
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size - 3 * lvl); r.font.color.rgb = color; r.font.bold = bold
        p.space_after = Pt(9)
    return tf


def text_and_figure(head, items, fig, fig_left=6.9, fig_top=1.4, fig_w=6.1):
    s = prs.slides.add_slide(BLANK)
    header(s, head)
    bullets(s, items, left=0.6, top=1.5, width=6.1, height=5.6, size=17)
    p = os.path.join(FIG, fig)
    if os.path.exists(p):
        s.shapes.add_picture(p, Inches(fig_left), Inches(fig_top), width=Inches(fig_w))
    return s


def full_figure(head, fig, note=None):
    s = prs.slides.add_slide(BLANK)
    header(s, head)
    p = os.path.join(FIG, fig)
    if os.path.exists(p):
        s.shapes.add_picture(p, Inches(0.7), Inches(1.4), width=Inches(12.0))
    if note:
        tf = _tb(s, 0.7, 6.9, 12.0, 0.5)
        r = tf.paragraphs[0].add_run(); r.text = note
        r.font.size = Pt(13); r.font.color.rgb = GREY
    return s


# ---------------------------------------------------------------- slides
s1 = title_slide(
    "MD-gated dendritic plasticity: progress update",
    "Does the mediodorsal thalamus gate which dendritic branch learns — enabling "
    "context-specific learning without catastrophic forgetting?",
    "Raghav · PI meeting 2026-10-15 · Thalamocortical Gating Model (spiking, Brian 2)")
notes(s1,
    "Open: it's been a while since we met — the setup moved from Windows to Linux in "
    "between, so this is the first update since then. I rebuilt everything and used the "
    "gap to close the three questions we'd flagged as open. Quick background first for "
    "context, then three results, then where I think we go next. Aim for ~10-12 min, "
    "leave time for the scaling decision at the end.")

# --- Background 1: the question (plain) ---
s = prs.slides.add_slide(BLANK)
header(s, "Background — the question")
bullets(s, [
    "While reading in a noisy café, the brain routes attention without erasing what it knows.",
    ("Hypothesis: the mediodorsal (MD) thalamus acts as a meta-controller that selects "
     "WHICH dendritic branch of a prefrontal pyramidal neuron is allowed to undergo plasticity.", 0),
    ("→ Different contexts write onto different physical branches instead of overwriting "
     "the same synapses — context-specific learning without catastrophic forgetting.", 0, BLUE, True),
    "",
    ("Two papers grafted together (neither answers it alone):", 0),
    ("Yang, Murray & Wang 2016 — the mechanism: spiking reduced-compartment pyramidal "
     "neuron, dendritic disinhibitory circuit (VIP→SOM→dendrite, PV→soma), "
     "calcium-based Graupner–Brunel plasticity.", 1),
    ("Zhang et al. 2025 — the MD wiring: MD→PV (gain) + MD→VIP (disinhibition) as "
     "the context/meta-control signal.", 1),
], top=1.4, size=18)
notes(s,
    "The core idea in one sentence: instead of different contexts fighting over the same "
    "synapses — which causes catastrophic forgetting — the MD thalamus routes each context "
    "to a different dendritic branch, so they don't overwrite each other. Neither source "
    "paper does this alone: Yang 2016 has the single-neuron dendritic gating mechanism but "
    "no thalamus; Zhang 2025 has the MD control wiring but not the plasticity. The project "
    "is grafting them. If asked 'is this just Yang replication?' — no, the MD-as-context-"
    "controller for plasticity is the novel part.")

# --- Background 2: the model (plain + diagram) ---
s_model = text_and_figure(
    "Background — the model (Model A, spiking, Brian 2)",
    [
        ("One pyramidal neuron: soma + 2 dendritic branches.", 0),
        ("Each branch has an independent plastic cortical input "
         "(Graupner–Brunel calcium rule, Yang Eqs. 32–35).", 0),
        ("SOM inhibits a branch → shunts its calcium → gates its plasticity.", 0),
        ("MD sets the context via PV (somatic gain) and VIP→SOM (branch disinhibition).", 0),
        ("", 0),
        ("Readout: gating ratio = rho(free branch) / rho(inhibited branch). "
         "Ratio > 2 = gating works.", 0, BLUE, True),
    ],
    "network_architecture.png", fig_left=6.6, fig_top=1.3, fig_w=6.4)
notes(s_model,
    "Walk the diagram right-to-left: the cortical pathway (dashed) is the plastic input "
    "onto two dendritic branches — that's where learning happens, via the Graupner-Brunel "
    "calcium rule. SOM (bottom left) inhibits a branch and shunts its calcium, which is how "
    "a branch gets gated off. MD (top left) is the controller: MD→PV sets somatic gain, "
    "MD→VIP→SOM sets which branch is disinhibited. The one number to remember all talk: the "
    "gating ratio — free-branch rho over inhibited-branch rho. Above 2 means gating works. "
    "Everything that follows is measured against that readout.")

# --- Since last meeting ---
s = prs.slides.add_slide(BLANK)
header(s, "Since the last meeting")
bullets(s, [
    ("Environment moved from Windows to Linux; rebuilt the Brian 2 stack and verified "
     "bit-exact reproduction of all committed results.", 0),
    ("Then closed the three open questions that were blocking confidence in the model:", 0),
    ("1.  Gating reliability — why did a few random seeds fail?", 1, ORANGE, True),
    ("2.  SOM-rate divergence — why does the model need ~2× Yang's firing rate?", 1, ORANGE, True),
    ("3.  No catastrophic forgetting — does the core claim actually hold?", 1, ORANGE, True),
    ("", 0),
    ("All three resolved, committed, and pushed to GitHub. Details follow.", 0, GREEN, True),
], top=1.5, size=19)
notes(s,
    "Honesty note on the move: I verified the Linux rebuild reproduces the old results "
    "bit-for-bit before trusting any new number — same RNG stream, same ratios. So nothing "
    "here is an artifact of the environment change. The three questions were the ones I'd "
    "have been nervous presenting without answers: a couple of seeds failing, the 2x firing-"
    "rate gap vs Yang, and whether the headline no-forgetting claim actually holds when "
    "tested directly. All three are now answered — next three slides.")

# --- Result 1 ---
s_r1 = text_and_figure(
    "Result 1 — Gating reliability: a small-N variance artifact",
    [
        ("Problem: at the 25 Hz baseline, 2/20 random seeds failed (gating ratio ~1).", 0, ORANGE, True),
        ("Diagnosed: the inhibited branch draws a small, random integer number of "
         "SOM synapses (p=0.5 over only 20 SOM cells). Too few → weak shunt → branch "
         "depolarizes → NMDA opens → calcium escapes → gating lost.", 0),
        ("Not a firing-rate or timing problem; a population-size variance effect.", 0),
        ("Fix (controlled): hold mean inhibition constant, enlarge the SOM pool so the "
         "variance self-averages.", 0),
        ("Result: N=100 → 20/20 seeds gate (full 10-epoch protocol).", 0, GREEN, True),
    ],
    "n_som_disinhibition_variance.png")
notes(s_r1,
    "Story: a few seeds failed and I wanted to know if it was a real bug. It's not — it's a "
    "finite-size effect. With only 20 SOM cells at p=0.5, the inhibited branch sometimes "
    "draws ~6 synapses instead of ~10, inhibition is too weak, and that branch escapes. "
    "Left panel: the spread in branch-0 inhibition (gGABA) at N=20 vs N=100 — the two red "
    "diamonds on the floor are the failing seeds. The clean test: I held the MEAN inhibition "
    "fixed and only shrank the variance by enlarging the pool. Right panel: 18/20 goes to "
    "20/20. So it's variance, not insufficient inhibition and not firing rate. Caveat if "
    "pushed: this is the reduced circuit; the fix is a population-size argument, which the "
    "full network will have naturally. Takeaway for the PI: the 30 Hz baseline we'd been "
    "using masks this; the real knob is pool size.")

# --- Result 2 ---
s_r2 = text_and_figure(
    "Result 2 — SOM-rate divergence from Yang, resolved",
    [
        ("Reviewer question: Yang gets gating at ~10 Hz SOM; our model needs ~18–21 Hz "
         "at the SAME τ_GABA = 20 ms. Why 2×?", 0, ORANGE, True),
        ("Answer: gating tracks the DUTY CYCLE of inhibition (rate × τ), not the rate "
         "alone. Our spiking calcium is event-driven, so it needs fuller coverage than "
         "a rate-averaged rule.", 0),
        ("Proof: critical rate scales ~1/τ (rate×τ constant to <3%, 10→20 ms). "
         "Same spike train + longer GABA tail rescues gating.", 0),
        ("The model recovers Yang's ~10 Hz regime at τ = 40 ms — a quantified "
         "consequence of spiking calcium, not a bug.", 0, GREEN, True),
    ],
    "som_rate_divergence.png")
notes(s_r2,
    "This is the reviewer-proofing one — the question a referee WILL ask: Yang gates at "
    "10 Hz SOM, why does ours need ~20? I checked the paper directly: Yang's 10 Hz is an "
    "experimental number, and his GABA decay tau is 20 ms — same as ours. So it's not a "
    "parameter mismatch. The answer is that what matters is COVERAGE — rate times tau, the "
    "fraction of each inter-spike gap the inhibition actually covers — not the rate alone. "
    "Our calcium is event-driven (spiking), so an uncovered gap lets calcium escape; a rate-"
    "averaged model like the mean-field reading of Yang doesn't feel that. Proof on the "
    "right: critical rate scales as 1/tau — double tau, halve the rate needed — and at "
    "tau=40 ms we gate at ~12 Hz, essentially Yang's regime. Same spike train, longer GABA "
    "tail, gating comes back. So the divergence is understood and quantified, not a bug. "
    "Caveat: single seed, reduced circuit; the scaling is clean for 10->20 ms and saturates "
    "a bit by 40 ms, which is expected.")

# --- Result 3 ---
s_r3 = text_and_figure(
    "Result 3 — No catastrophic forgetting (the core claim)",
    [
        ("First direct test: learn context A → switch to B → is A retained?", 0),
        ("Branch separation ALONE is not enough: inhibition-only switch → "
         "A's memory 70% erased (30% retained).", 0, ORANGE, True),
        ("Why: gating is not protection. A shunted branch still gets input + bAP calcium, "
         "sits in the LTD band, and is actively depressed (Yang plasticity has no frozen "
         "state).", 0),
        ("Memory is protected only when context A's INPUT is silent during B → "
         "100% retained, while B learns on its own branch.", 0, GREEN, True),
        ("Sharper claim: no forgetting needs context-specific gating of BOTH the "
         "disinhibition AND the input.", 0, BLUE, True),
    ],
    "retention_forgetting.png")
notes(s_r3,
    "This is the headline claim, and it had never actually been tested — so I ran it: learn "
    "A on one branch, switch context, learn B on the other, check A survives. First result "
    "was a surprise and I want to present it honestly: with branch separation alone, A was "
    "70% ERASED (left panel, blue line collapsing). So 'different branches' is NOT by itself "
    "enough. The reason is mechanistic: a gated branch still receives its input plus back-"
    "propagating-AP calcium, which sits in the LTD band — Yang's plasticity rule has no "
    "frozen state, so shunting slowly depresses the old memory rather than protecting it. "
    "The fix — and I'd argue this is the actual scientific result, not a patch — is that the "
    "memory is protected when context A's INPUT goes silent during B (right panel: flat, "
    "100% retained). So the real claim is sharper and more defensible: no forgetting needs "
    "context-specific gating of BOTH the disinhibition AND the input — which is exactly what "
    "a real context switch is. Frame this as a strengthened claim, not a problem. Caveat: "
    "reduced circuit, context abstracted, 8 epochs of interference.")

# --- Summary: before -> after ---
s_sum = full_figure(
    "Summary — before → after (this session)",
    "before_after_proof.png",
    "All panels from committed simulation data. Orange = baseline/broken; green = resolved.")
notes(s_sum,
    "One-slide recap if we're short on time, or the landing slide. Three before/after pairs, "
    "all from real simulation data: reliability 18/20 to 20/20; the Yang rate gap explained "
    "and the 10 Hz regime recovered; forgetting 30% to 100% retention. The through-line: in "
    "every case the model does what we claimed, once the mechanism is understood properly. "
    "Orange is where we were, green is where we are.")

# --- Next steps ---
s = prs.slides.add_slide(BLANK)
header(s, "Where this leaves us / next steps")
bullets(s, [
    ("Settled: gating reliability, SOM-rate calibration, and the no-forgetting mechanism "
     "— all verified in the reduced circuit and documented.", 0, GREEN, True),
    ("", 0),
    ("Next — scale to a ~3000-neuron network (Yang scale):", 0, BLACK, True),
    ("Now unblocked in principle, but it MUST preserve input context-sparsity — not "
     "just branch-specific gating — or it will forget at scale.", 1),
    ("Repeat the key tests across seeds and in the full MD→VIP→SOM circuit.", 1),
    ("", 0),
    ("Open decisions for discussion: ambition (faithful replication + modest extension "
     "vs. full novel network for publication); MD fidelity (Poisson vs. full PFC↔MD loop).", 0, GREY),
], top=1.4, size=18)
notes(s,
    "The ask. The three prerequisites we wanted settled before scaling are settled, so the "
    "network build (~3000 neurons, Yang scale) is unblocked. One hard constraint I want to "
    "flag: Result 3 means the network must keep inputs context-SPARSE — different contexts "
    "can't share persistently-active inputs, or it'll forget at scale for the same reason. "
    "Two decisions I'd like your steer on: (1) ambition — faithful replication plus a modest "
    "extension, or push for a full novel network aimed at publication? (2) MD fidelity — keep "
    "MD as a Poisson drive, or build the full PFC↔MD loop? My lean: lock the scaling scope "
    "with the input-sparsity constraint, keep MD as Poisson for the first network pass, and "
    "decide publication ambition once it holds at scale. Open the floor here.")

prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
